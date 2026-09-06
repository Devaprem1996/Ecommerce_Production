import csv
import json
import os

# Ultra-brief, crisp requirements matrix focusing strictly on the single Responsible Person (Dev, GM, or Client)
raci_data = [
    {
        "ID": "REQ-01",
        "Category": "Catalog Data",
        "Requirement": "Finalize 85 product catalog items (EN & TA names, pricing & variants)",
        "Responsible_Person": "Dev",
        "Priority": "Critical",
        "Status": "Completed",
        "Action_Required": "Seeded into backend/prisma/seed.ts dataset."
    },
    {
        "ID": "REQ-02",
        "Category": "Photography",
        "Requirement": "Collect product photos & remove background clutter (White background)",
        "Responsible_Person": "GM",
        "Priority": "High",
        "Status": "In Progress",
        "Action_Required": "Photograph items; remove backgrounds with PhotoRoom / Remove.bg."
    },
    {
        "ID": "REQ-03",
        "Category": "Image Storage",
        "Requirement": "Create Cloudinary account & setup product image upload preset",
        "Responsible_Person": "Dev",
        "Priority": "High",
        "Status": "Live Ready",
        "Action_Required": "Cloudinary credentials & upload middleware configured."
    },
    {
        "ID": "REQ-04",
        "Category": "Hosting Plan",
        "Requirement": "Purchase Hostinger VPS plan (Ubuntu 22.04 LTS - India Location)",
        "Responsible_Person": "Client",
        "Priority": "Critical",
        "Status": "Pending",
        "Action_Required": "Client to buy Hostinger KVM VPS plan."
    },
    {
        "ID": "REQ-05",
        "Category": "Database Setup",
        "Requirement": "Create PostgreSQL database on Hostinger VPS & link DATABASE_URL",
        "Responsible_Person": "Dev",
        "Priority": "Critical",
        "Status": "Live Ready",
        "Action_Required": "PostgreSQL setup script & Prisma config prepared."
    },
    {
        "ID": "REQ-06",
        "Category": "Domain Purchase",
        "Requirement": "Purchase domain name (.in or .com) via Hostinger / Namecheap",
        "Responsible_Person": "Client",
        "Priority": "Critical",
        "Status": "Pending",
        "Action_Required": "Client to register domain name."
    },
    {
        "ID": "REQ-07",
        "Category": "DNS & SSL",
        "Requirement": "Point Domain DNS A-Records to Hostinger IP & issue free SSL cert",
        "Responsible_Person": "Dev",
        "Priority": "High",
        "Status": "Ready",
        "Action_Required": "Configure Nginx reverse proxy & Certbot SSL."
    },
    {
        "ID": "REQ-08",
        "Category": "Live Deploy",
        "Requirement": "Run database migrations & seed 85 products to production database",
        "Responsible_Person": "Dev",
        "Priority": "Critical",
        "Status": "Ready",
        "Action_Required": "Execute npx prisma db seed on live VPS."
    },
    {
        "ID": "REQ-09",
        "Category": "Payment Gateway",
        "Requirement": "Integrate Razorpay / Cashfree UPI & Card online payments",
        "Responsible_Person": "Dev",
        "Priority": "Medium",
        "Status": "Phase 2 (Hold)",
        "Action_Required": "On Hold for Phase 2 as instructed."
    },
    {
        "ID": "REQ-10",
        "Category": "Shipping Gateway",
        "Requirement": "Integrate Shiprocket courier dispatch API & tracking",
        "Responsible_Person": "Dev",
        "Priority": "Medium",
        "Status": "Phase 2 (Hold)",
        "Action_Required": "On Hold for Phase 2 as instructed."
    },
    {
        "ID": "REQ-11",
        "Category": "Local Pincodes",
        "Requirement": "Set up South India pincodes & flat delivery rates",
        "Responsible_Person": "Dev",
        "Priority": "High",
        "Status": "Completed",
        "Action_Required": "Seeded Chennai, Madurai, Kovai, Trichy, Blr pincodes."
    },
    {
        "ID": "REQ-12",
        "Category": "Admin Security",
        "Requirement": "Setup admin credentials (admin@yathu.com) & JWT authentication",
        "Responsible_Person": "Dev",
        "Priority": "Critical",
        "Status": "Live Ready",
        "Action_Required": "CMS routes secured with JWT role authorization."
    },
    {
        "ID": "REQ-13",
        "Category": "Admin SOP",
        "Requirement": "Prepare guide for adding/editing products & updating prices",
        "Responsible_Person": "Dev",
        "Priority": "High",
        "Status": "Completed",
        "Action_Required": "Documented in MASTER_SETUP_AND_PRODUCT_OPERATIONS_GUIDE.md."
    },
    {
        "ID": "REQ-14",
        "Category": "DB Backups",
        "Requirement": "Setup automated daily PostgreSQL database backups",
        "Responsible_Person": "Dev",
        "Priority": "High",
        "Status": "Ready",
        "Action_Required": "pg_dump automated cron script created."
    }
]

def generate_csv(filename):
    fieldnames = ["ID", "Category", "Requirement", "Responsible_Person", "Priority", "Status", "Action_Required"]
    try:
        with open(filename, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for row in raci_data:
                writer.writerow(row)
        print(f"Generated CSV: {filename}")
    except PermissionError:
        alt_name = filename.replace(".csv", "_Updated.csv")
        with open(alt_name, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for row in raci_data:
                writer.writerow(row)
        print(f"File locked. Generated alternative CSV: {alt_name}")

def generate_xlsx(filename):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Requirements Matrix"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")

        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB')
        )

        headers = ["ID", "Category", "Requirement", "Responsible Person", "Priority", "Status", "Action Required"]
        ws.append(headers)

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        for row_idx, row_data in enumerate(raci_data, start=2):
            ws.append([
                row_data["ID"],
                row_data["Category"],
                row_data["Requirement"],
                row_data["Responsible_Person"],
                row_data["Priority"],
                row_data["Status"],
                row_data["Action_Required"]
            ])

            ws.cell(row=row_idx, column=1).alignment = align_center
            ws.cell(row=row_idx, column=2).alignment = align_left
            ws.cell(row=row_idx, column=3).alignment = align_left
            
            resp_cell = ws.cell(row=row_idx, column=4)
            resp_cell.alignment = align_center
            resp_cell.font = Font(bold=True)
            if row_data["Responsible_Person"] == "Dev":
                resp_cell.fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
            elif row_data["Responsible_Person"] == "GM":
                resp_cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
            elif row_data["Responsible_Person"] == "Client":
                resp_cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

            ws.cell(row=row_idx, column=5).alignment = align_center
            ws.cell(row=row_idx, column=6).alignment = align_center
            ws.cell(row=row_idx, column=7).alignment = align_left

            for c in range(1, 8):
                ws.cell(row=row_idx, column=c).border = thin_border

        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 60
        ws.column_dimensions['D'].width = 22
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 16
        ws.column_dimensions['G'].width = 50

        try:
            wb.save(filename)
            print(f"Generated Crisp XLSX: {filename}")
        except PermissionError:
            alt_xlsx = filename.replace(".xlsx", "_Updated.xlsx")
            wb.save(alt_xlsx)
            print(f"File locked. Generated alternative XLSX: {alt_xlsx}")
    except ImportError:
        print("openpyxl not installed.")

if __name__ == "__main__":
    generate_csv("Requirements_RACI_Matrix.csv")
    generate_xlsx("Requirements_RACI_Matrix.xlsx")
